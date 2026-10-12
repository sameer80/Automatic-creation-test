"""Data-Driven Driver — run the same steps once per Excel/CSV/JSON data row."""

from __future__ import annotations

import csv
import json
import re
from copy import deepcopy
from dataclasses import asdict
from pathlib import Path
from typing import Any

from dynamic.discovery import ScenarioDef, StepDef

ROOT = Path(__file__).resolve().parent.parent
GENERATED_DIR = ROOT / "generated"
DEFAULT_TESTDATA = GENERATED_DIR / "TestData.xlsx"
AI_TESTDATA = GENERATED_DIR / "AI_TestData.xlsx"
DRIVERS_CONFIG = ROOT / "config" / "drivers.json"

_PLACEHOLDER_RE = re.compile(r"\{\{\s*([^}]+)\s*\}\}|\{([A-Za-z_][A-Za-z0-9_]*)\}")


def load_drivers_config() -> dict[str, Any]:
    defaults = {
        "data_driven_enabled": True,
        "hybrid_enabled": True,
        "default_mode": "linear",
        "testdata_file": "generated/TestData.xlsx",
        "resolve_placeholders_at_runtime": True,
    }
    if not DRIVERS_CONFIG.is_file():
        return defaults
    try:
        data = json.loads(DRIVERS_CONFIG.read_text(encoding="utf-8-sig"))
        if isinstance(data, dict):
            return {**defaults, **data}
    except Exception:  # noqa: BLE001
        pass
    return defaults


def is_data_driven_enabled() -> bool:
    return bool(load_drivers_config().get("data_driven_enabled", True))


def resolve_placeholders(value: str, data: dict[str, Any] | None) -> str:
    """Replace {{Column}} or {Column} with values from the data row / vars."""
    if not value or not data:
        return value or ""
    text = str(value)

    def repl(match: re.Match[str]) -> str:
        key = (match.group(1) or match.group(2) or "").strip()
        if not key:
            return match.group(0)
        if key in data:
            return str(data[key])
        for k, v in data.items():
            if str(k).lower() == key.lower():
                return str(v)
        return match.group(0)

    return _PLACEHOLDER_RE.sub(repl, text)


def apply_data_to_step(step: StepDef, data: dict[str, Any]) -> StepDef:
    """Return a copy of the step with placeholders resolved from data."""
    return StepDef(
        step_no=step.step_no,
        action=step.action,
        object_name=resolve_placeholders(step.object_name, data),
        input_value=resolve_placeholders(step.input_value, data),
        locator_by=step.locator_by,
        locator_value=resolve_placeholders(step.locator_value, data),
        expected=resolve_placeholders(step.expected, data),
    )


def load_testdata(
    path: Path | str | None = None,
    sheet: str | None = None,
) -> list[dict[str, Any]]:
    """Load rows from .xlsx / .csv / .json. Returns list of column→value dicts."""
    file_path = Path(path) if path else _default_testdata_path()
    if not file_path.is_file():
        return []

    suffix = file_path.suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        return _load_excel(file_path, sheet)
    if suffix == ".csv":
        return _load_csv(file_path)
    if suffix == ".json":
        return _load_json(file_path, sheet)
    return []


def _default_testdata_path() -> Path:
    cfg = load_drivers_config()
    rel = str(cfg.get("testdata_file") or "generated/TestData.xlsx")
    candidate = ROOT / rel
    if candidate.is_file():
        return candidate
    if AI_TESTDATA.is_file():
        return AI_TESTDATA
    return DEFAULT_TESTDATA


def _load_excel(path: Path, sheet: str | None) -> list[dict[str, Any]]:
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True, read_only=True)
    try:
        if sheet and sheet in wb.sheetnames:
            ws = wb[sheet]
        else:
            ws = wb.active
        rows_iter = ws.iter_rows(values_only=True)
        try:
            headers = next(rows_iter)
        except StopIteration:
            return []
        headers = [str(h).strip() if h is not None else f"col{i}" for i, h in enumerate(headers)]
        out: list[dict[str, Any]] = []
        for raw in rows_iter:
            if raw is None or all(c is None or str(c).strip() == "" for c in raw):
                continue
            row = {
                headers[i]: ("" if raw[i] is None else raw[i])
                for i in range(min(len(headers), len(raw)))
                if headers[i]
            }
            # Skip disabled rows
            enabled = str(row.get("Enabled", row.get("enabled", "Y"))).strip().upper()
            if enabled in {"N", "NO", "FALSE", "0"}:
                continue
            out.append(row)
        return out
    finally:
        wb.close()


def _load_csv(path: Path) -> list[dict[str, Any]]:
    with path.open(encoding="utf-8-sig", newline="") as fh:
        reader = csv.DictReader(fh)
        rows = []
        for row in reader:
            enabled = str(row.get("Enabled", row.get("enabled", "Y"))).strip().upper()
            if enabled in {"N", "NO", "FALSE", "0"}:
                continue
            rows.append({k: (v if v is not None else "") for k, v in row.items() if k})
        return rows


def _load_json(path: Path, sheet: str | None) -> list[dict[str, Any]]:
    raw = json.loads(path.read_text(encoding="utf-8-sig"))
    if isinstance(raw, list):
        return [r for r in raw if isinstance(r, dict)]
    if isinstance(raw, dict):
        if sheet and isinstance(raw.get(sheet), list):
            return [r for r in raw[sheet] if isinstance(r, dict)]
        for key in ("rows", "data", "sheets"):
            if isinstance(raw.get(key), list):
                return [r for r in raw[key] if isinstance(r, dict)]
            if isinstance(raw.get(key), dict) and sheet and isinstance(raw[key].get(sheet), list):
                return [r for r in raw[key][sheet] if isinstance(r, dict)]
    return []


def expand_scenario_with_rows(
    scenario: ScenarioDef,
    rows: list[dict[str, Any]],
    *,
    resolve_now: bool = True,
) -> list[ScenarioDef]:
    """Clone scenario once per data row (Data-Driven Driver)."""
    if not rows:
        return [scenario]

    expanded: list[ScenarioDef] = []
    for idx, row in enumerate(rows, start=1):
        row_label = str(row.get("TestCase") or row.get("Name") or row.get("Scenario") or idx)
        steps: list[StepDef] = []
        for st in scenario.steps:
            if resolve_now:
                steps.append(apply_data_to_step(st, row))
            else:
                steps.append(
                    StepDef(
                        step_no=st.step_no,
                        action=st.action,
                        object_name=st.object_name,
                        input_value=st.input_value,
                        locator_by=st.locator_by,
                        locator_value=st.locator_value,
                        expected=st.expected,
                    )
                )
        for i, st in enumerate(steps, start=1):
            st.step_no = i
        sc = ScenarioDef(
            id=f"{scenario.id}_D{idx}",
            type=scenario.type if scenario.type != "hybrid" else "automation",
            title=f"{scenario.title} [Data {idx}: {row_label}]",
            steps=steps,
            module=scenario.module,
        )
        # Attach row for runtime resolution / reporting
        sc._data_row = row  # type: ignore[attr-defined]
        sc._driver = "data_driven"  # type: ignore[attr-defined]
        expanded.append(sc)
    return expanded


def expand_scenarios_data_driven(
    scenarios: list[ScenarioDef],
    *,
    data_path: Path | str | None = None,
    sheet: str | None = None,
    rows: list[dict[str, Any]] | None = None,
) -> list[ScenarioDef]:
    """Expand all scenarios that need data rows."""
    if not is_data_driven_enabled():
        return list(scenarios)

    data_rows = list(rows) if rows is not None else load_testdata(data_path, sheet)
    out: list[ScenarioDef] = []
    for sc in scenarios:
        embedded = getattr(sc, "_data_rows", None)
        if isinstance(embedded, list) and embedded:
            out.extend(expand_scenario_with_rows(sc, embedded))
            continue
        # Expand when placeholders exist OR type is data_driven
        needs_data = (
            (sc.type or "").lower() in {"data_driven", "datadriven", "ddt"}
            or any("{{" in (st.input_value or "") or "{{" in (st.expected or "") for st in sc.steps)
        )
        if needs_data and data_rows:
            out.extend(expand_scenario_with_rows(sc, data_rows))
        else:
            out.append(sc)
    return out


def scenario_from_dict(raw: dict[str, Any]) -> ScenarioDef:
    steps = []
    for index, st in enumerate(raw.get("steps") or [], start=1):
        if not isinstance(st, dict):
            continue
        steps.append(
            StepDef(
                step_no=int(st.get("step_no") or index),
                action=str(st.get("action") or "Verify"),
                object_name=str(st.get("object_name") or ""),
                input_value=str(st.get("input_value") or ""),
                locator_by=str(st.get("locator_by") or ""),
                locator_value=str(st.get("locator_value") or ""),
                expected=str(st.get("expected") or ""),
            )
        )
    sc = ScenarioDef(
        id=str(raw.get("id") or "TC_DD"),
        type=str(raw.get("type") or "data_driven"),
        title=str(raw.get("title") or raw.get("name") or "Data driven"),
        steps=steps,
        module=str(raw.get("module") or ""),
    )
    if raw.get("data_rows"):
        sc._data_rows = list(raw["data_rows"])  # type: ignore[attr-defined]
    return sc


def write_sample_testdata(path: Path | None = None) -> Path:
    """Create a starter TestData.xlsx for Data-Driven / Hybrid runs."""
    from openpyxl import Workbook

    GENERATED_DIR.mkdir(parents=True, exist_ok=True)
    out = path or DEFAULT_TESTDATA
    wb = Workbook()
    ws = wb.active
    ws.title = "Login"
    ws.append(["Enabled", "TestCase", "Username", "Password", "Expected"])
    ws.append(["Y", "Valid login", "Admin", "admin123", "Dashboard"])
    ws.append(["Y", "Invalid password", "Admin", "wrong", "Invalid"])
    ws.append(["N", "Skipped row", "x", "y", "n/a"])

    ws2 = wb.create_sheet("Form")
    ws2.append(["Enabled", "TestCase", "Name", "Email", "Amount", "Expected"])
    ws2.append(["Y", "Happy path", "Sameer", "sameer@example.com", "100", "Success"])
    ws2.append(["Y", "Blank name", "", "a@b.com", "50", "Required"])

    wb.save(out)
    return out


def merge_data_context(context: dict[str, Any] | None, row: dict[str, Any] | None) -> dict[str, Any]:
    """Merge data row into runner context vars (row wins for same keys)."""
    ctx = context if context is not None else {}
    vars_store = ctx.setdefault("vars", {})
    if row:
        ctx["data_row"] = row
        for key, value in row.items():
            vars_store[str(key)] = value
    return ctx


def get_data_row_for_scenario(scenario: ScenarioDef) -> dict[str, Any]:
    return dict(getattr(scenario, "_data_row", None) or {})
